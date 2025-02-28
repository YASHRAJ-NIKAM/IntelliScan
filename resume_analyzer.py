import streamlit as st
import pdfplumber
import docx
import pandas as pd
import spacy
import os
from fpdf import FPDF

# Load NLP model
nlp = spacy.load("en_core_web_sm")

# Extract text from resume (PDF/DOCX)
def extract_text(uploaded_file):
    file_extension = os.path.splitext(uploaded_file.name)[-1].lower()
    if file_extension == ".pdf":
        with pdfplumber.open(uploaded_file) as pdf:
            return "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])
    elif file_extension == ".docx":
        doc = docx.Document(uploaded_file)
        return "\n".join([para.text for para in doc.paragraphs])
    return None

# AI-Based Skill Extraction
def extract_skills_nlp(text):
    doc = nlp(text)
    extracted_skills = {token.text for token in doc if token.pos_ in ["NOUN", "PROPN"]}
    return extracted_skills

# Resume Formatting Check
def check_resume_format(resume_text):
    sections = {
        "Education": ["education", "degree", "bachelor", "master", "university"],
        "Experience": ["experience", "internship", "company", "work"],
        "Skills": ["skills", "technologies", "expertise"],
        "Certifications": ["certification", "course", "training"],
        "Projects": ["projects", "portfolio", "github"]
    }
    return [section for section, keywords in sections.items() if not any(keyword in resume_text.lower() for keyword in keywords)]

# Generate PDF Report
def generate_pdf(report_data):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(200, 10, "AI Resume Analyzer Report", ln=True, align="C")
    
    pdf.set_font("Arial", "", 12)
    for key, value in report_data.items():
        pdf.cell(200, 10, f"{key}: {value}", ln=True)
    
    return pdf.output(dest="S").encode("latin1")

# Streamlit UI Setup
st.title("📄 AI Resume Analyzer")
uploaded_file = st.file_uploader("Upload your resume (PDF/DOCX)", type=["pdf", "docx"])
if uploaded_file:
    resume_text = extract_text(uploaded_file)
    if resume_text:
        extracted_skills = extract_skills_nlp(resume_text)
        missing_sections = check_resume_format(resume_text)
        
        st.subheader("Extracted Skills")
        st.write(extracted_skills)
        
        st.subheader("Missing Resume Sections")
        st.write(missing_sections if missing_sections else "None")
        
        report_data = {
            "Extracted Skills": ", ".join(extracted_skills),
            "Missing Sections": ", ".join(missing_sections) if missing_sections else "None"
        }
        pdf_bytes = generate_pdf(report_data)
        st.download_button(label="📥 Download Report (PDF)", data=pdf_bytes, file_name="resume_analysis.pdf", mime="application/pdf")
